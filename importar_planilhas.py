from __future__ import annotations

from datetime import datetime
from pathlib import Path
import re
import unicodedata

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.database import Base, engine
from app.models import (
    Acesso,
    Departamento,
    Funcionario,
    FuncionarioPerfil,
    Perfil,
    PerfilAcesso,
)

ROOT = Path(__file__).resolve().parent
IMPORTS = ROOT / "app" / "imports"
BASE_XLSX = IMPORTS / "base_perfis.xlsx"
PERFIS_DIR = IMPORTS / "perfis"


def clean(value) -> str:
    return " ".join(str(value or "").replace("\xa0", " ").split()).strip()


def key(value) -> str:
    return (
        unicodedata.normalize("NFKD", clean(value))
        .encode("ascii", "ignore")
        .decode()
        .casefold()
    )


def bool_value(value, default=False) -> bool:
    text = key(value)
    if text in {"sim", "s", "1", "true", "ativo", "vinculado", "permitido"}:
        return True
    if text in {"nao", "n", "0", "false", "inativo", "nao vinculado", "negado"}:
        return False
    return default


def parse_date(value):
    if isinstance(value, datetime):
        return value
    text = clean(value)
    for fmt in ("%d/%m/%Y %H:%M:%S", "%d/%m/%Y", "%Y-%m-%d %H:%M:%S"):
        try:
            return datetime.strptime(text, fmt)
        except ValueError:
            pass
    return None


def headers(ws) -> dict[str, int]:
    return {key(cell.value): cell.column - 1 for cell in ws[1]}


def value(row, mapping, name) -> str:
    position = mapping.get(key(name))
    return clean(row[position]) if position is not None and position < len(row) else ""


def get_or_create_department(session: Session, name: str | None):
    name = clean(name)
    if not name:
        return None
    item = session.scalar(select(Departamento).where(Departamento.nome == name))
    if item is None:
        item = Departamento(nome=name, ativo=True)
        session.add(item)
        session.flush()
    return item


def get_or_create_profile(session: Session, name: str):
    item = session.scalar(select(Perfil).where(Perfil.nome == name))
    if item is None:
        item = Perfil(nome=name, situacao="Pendente")
        session.add(item)
        session.flush()
    return item


def get_or_create_employee(session: Session, login: str, name: str = ""):
    item = session.scalar(select(Funcionario).where(Funcionario.login == login))
    if item is None:
        item = Funcionario(login=login, nome=name or login, ativo=True)
        session.add(item)
        session.flush()
    elif name:
        item.nome = name
    return item


def link_employee_profile(session: Session, employee, profile, origin="Importacao", collected=None):
    link = session.scalar(
        select(FuncionarioPerfil).where(
            FuncionarioPerfil.funcionario_id == employee.id,
            FuncionarioPerfil.perfil_id == profile.id,
        )
    )
    if link is None:
        link = FuncionarioPerfil(
            funcionario_id=employee.id,
            perfil_id=profile.id,
            origem=origin or "Importacao",
        )
        session.add(link)
    else:
        link.origem = origin or link.origem
    if collected:
        link.coletado_em = collected
    return link


def get_or_create_access(session: Session, path: str, screen: str, button: str):
    access_key = "|||".join((path, screen, button))
    item = session.scalar(select(Acesso).where(Acesso.chave_acesso == access_key))
    if item is None:
        item = Acesso(
            caminho_completo=path,
            tela_final=screen,
            botao=button,
            chave_acesso=access_key,
        )
        session.add(item)
        session.flush()
    return item


def set_profile_access(session: Session, profile, access, status: str, collected=None):
    item = session.scalar(
        select(PerfilAcesso).where(
            PerfilAcesso.perfil_id == profile.id,
            PerfilAcesso.acesso_id == access.id,
        )
    )
    normalized = clean(status) or "Nao identificado"
    if item is None:
        item = PerfilAcesso(
            perfil_id=profile.id,
            acesso_id=access.id,
            status=normalized,
        )
        session.add(item)
    else:
        item.status = normalized
    if collected:
        item.coletado_em = collected
    return item


def import_base(session: Session, path: Path, counters: dict):
    if not path.exists():
        print(f"Base consolidada nao encontrada: {path}")
        return

    wb = load_workbook(path, read_only=True, data_only=True)

    if "Funcionarios" in wb.sheetnames:
        ws = wb["Funcionarios"]
        h = headers(ws)
        for row in ws.iter_rows(min_row=2, values_only=True):
            login = value(row, h, "Login")
            if not login:
                continue
            employee = get_or_create_employee(session, login, value(row, h, "Nome"))
            department = get_or_create_department(session, value(row, h, "Departamento"))
            employee.departamento_id = department.id if department else None
            employee.incluir_revisao = bool_value(value(row, h, "Incluir na revisao"), False)
            employee.ativo = bool_value(value(row, h, "Ativo"), True)
            employee.observacao = value(row, h, "Observacao") or None
            counters["funcionarios"] += 1

    if "Perfis" in wb.sheetnames:
        ws = wb["Perfis"]
        h = headers(ws)
        for row in ws.iter_rows(min_row=2, values_only=True):
            name = value(row, h, "Perfil")
            if not name:
                continue
            profile = get_or_create_profile(session, name)
            profile.situacao = value(row, h, "Situacao") or profile.situacao
            profile.ultima_atualizacao = parse_date(value(row, h, "Ultima atualizacao"))
            counters["perfis"] += 1

    if "Funcionario_Perfis" in wb.sheetnames:
        ws = wb["Funcionario_Perfis"]
        h = headers(ws)
        for row in ws.iter_rows(min_row=2, values_only=True):
            login = value(row, h, "Login")
            profile_name = value(row, h, "Perfil")
            if not login or not profile_name:
                continue
            employee = get_or_create_employee(session, login, value(row, h, "Nome"))
            profile = get_or_create_profile(session, profile_name)
            link_employee_profile(
                session,
                employee,
                profile,
                value(row, h, "Origem") or "Importacao base",
                parse_date(value(row, h, "Data")),
            )
            counters["vinculos"] += 1

    if "Permissoes" in wb.sheetnames:
        ws = wb["Permissoes"]
        h = headers(ws)
        for row in ws.iter_rows(min_row=2, values_only=True):
            profile_name = value(row, h, "Perfil")
            path_value = value(row, h, "Caminho completo")
            screen = value(row, h, "Tela final")
            button = value(row, h, "Botao")
            if not profile_name or not path_value or not button:
                continue
            profile = get_or_create_profile(session, profile_name)
            access = get_or_create_access(session, path_value, screen, button)
            set_profile_access(
                session,
                profile,
                access,
                value(row, h, "Status"),
                parse_date(value(row, h, "Data")),
            )
            counters["permissoes"] += 1

    wb.close()
    print(f"Base importada: {path.name}")


def profile_name_from_file(path: Path) -> str:
    name = re.sub(r"^perfil_", "", path.stem, flags=re.IGNORECASE)
    name = re.sub(r"_\d{8}_\d{6}$", "", name)
    return clean(name.replace("_", " "))


def import_profile_file(session: Session, path: Path, counters: dict):
    wb = load_workbook(path, read_only=True, data_only=True)
    profile_name = ""

    access_sheet = "Acessos" if "Acessos" in wb.sheetnames else None
    if access_sheet:
        ws = wb[access_sheet]
        h = headers(ws)
        for row in ws.iter_rows(min_row=2, values_only=True):
            row_profile = value(row, h, "Perfil")
            if row_profile:
                profile_name = row_profile
                break

    profile_name = profile_name or profile_name_from_file(path)
    if not profile_name:
        wb.close()
        raise ValueError("Nome do perfil nao identificado")

    profile = get_or_create_profile(session, profile_name)
    access_count = 0
    user_count = 0

    if access_sheet:
        ws = wb[access_sheet]
        h = headers(ws)
        for row in ws.iter_rows(min_row=2, values_only=True):
            path_value = value(row, h, "Caminho completo")
            screen = value(row, h, "Tela final")
            button = value(row, h, "Botao")
            if not path_value or not button:
                continue
            access = get_or_create_access(session, path_value, screen, button)
            set_profile_access(
                session,
                profile,
                access,
                value(row, h, "Status"),
                parse_date(value(row, h, "Data")),
            )
            access_count += 1
            counters["permissoes"] += 1

    if "Usuarios" in wb.sheetnames:
        ws = wb["Usuarios"]
        h = headers(ws)
        for row in ws.iter_rows(min_row=2, values_only=True):
            login = value(row, h, "Identificador") or value(row, h, "Login")
            if not login:
                continue
            employee = get_or_create_employee(session, login, value(row, h, "Nome"))
            state = key(value(row, h, "Situacao"))
            if state == "vinculado":
                link_employee_profile(
                    session,
                    employee,
                    profile,
                    "Importacao planilha de perfil",
                    parse_date(value(row, h, "Data")),
                )
                counters["vinculos"] += 1
            user_count += 1

    profile.situacao = "Completo" if access_count and user_count else "Pendente"
    profile.ultima_atualizacao = datetime.now()
    wb.close()
    counters["arquivos_perfil"] += 1
    print(f"Perfil importado: {profile_name} | acessos={access_count} | usuarios={user_count}")


def main():
    Base.metadata.create_all(bind=engine)
    IMPORTS.mkdir(parents=True, exist_ok=True)
    PERFIS_DIR.mkdir(parents=True, exist_ok=True)

    counters = {
        "funcionarios": 0,
        "perfis": 0,
        "vinculos": 0,
        "permissoes": 0,
        "arquivos_perfil": 0,
        "erros": 0,
    }

    with Session(engine) as session:
        try:
            import_base(session, BASE_XLSX, counters)
            session.commit()
        except Exception as error:
            session.rollback()
            counters["erros"] += 1
            print(f"Erro na base consolidada: {error}")

        for file_path in sorted(PERFIS_DIR.glob("*.xlsx")):
            try:
                import_profile_file(session, file_path, counters)
                session.commit()
            except Exception as error:
                session.rollback()
                counters["erros"] += 1
                print(f"Erro em {file_path.name}: {error}")

    print("\nImportacao concluida.")
    for name, total in counters.items():
        print(f"{name}: {total}")


if __name__ == "__main__":
    main()
