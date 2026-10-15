from __future__ import annotations

import hashlib
import re
import unicodedata
from pathlib import Path
from collections import defaultdict
from io import BytesIO

from openpyxl import load_workbook
from sqlalchemy import select

from app.models import Acesso, Funcionario, FuncionarioPerfil, Perfil, PerfilAcesso


ALLOWED = {"v", "x", "✔", "✓", "1", "sim", "s", "permitido", "true"}
DENIED = {"✘", "×", "0", "nao", "não", "n", "negado", "false"}


def text(value):
    return "" if value is None else str(value).strip()


def key(value):
    normalized = unicodedata.normalize("NFKD", text(value))
    normalized = normalized.encode("ascii", "ignore").decode().lower()
    return re.sub(r"[^a-z0-9]+", " ", normalized).strip()


def compact_key(value):
    return key(value).replace(" ", "")


def normalize_login(value):
    return text(value).lower().replace(" ", "")


def workbook(data: bytes):
    return load_workbook(BytesIO(data), read_only=True, data_only=True)


def workbook_info(data: bytes):
    wb = workbook(data)
    return wb, [
        {"nome": ws.title, "linhas": ws.max_row, "colunas": ws.max_column}
        for ws in wb.worksheets
    ]


def find_sheet(wb, *names):
    wanted = {key(name) for name in names}
    for ws in wb.worksheets:
        if key(ws.title) in wanted:
            return ws
    return None


def header_map(ws, row_number=1):
    values = next(
        ws.iter_rows(
            min_row=row_number,
            max_row=row_number,
            values_only=True,
        ),
        (),
    )
    return {key(value): index for index, value in enumerate(values) if text(value)}


def row_value(values, headers, *names):
    for name in names:
        wanted = key(name)
        index = headers.get(wanted)
        if index is not None and index < len(values) and text(values[index]):
            return values[index]
    return None


def iter_records(ws, header_row=1):
    headers = header_map(ws, header_row)
    for row_number, values in enumerate(
        ws.iter_rows(min_row=header_row + 1, values_only=True),
        start=header_row + 1,
    ):
        if any(text(value) for value in values):
            yield row_number, values, headers


def _status(value):
    normalized = key(value)
    if normalized in {key(item) for item in DENIED}:
        return "Negado"
    if normalized in {key(item) for item in ALLOWED}:
        return "Permitido"
    if normalized == "negado":
        return "Negado"
    return "Permitido"


def _access_identity(path, screen, button):
    return compact_key(path), compact_key(screen), compact_key(button)


def _profile_identity(name):
    return compact_key(name)


def _employee_name_identity(name):
    return compact_key(name)


def _employee_login_identity(login):
    return normalize_login(login)


def _natural_access_key(path, screen, button):
    return "|".join([text(path), text(screen), text(button)])


def preview_files(consulta_data: bytes, mapa_data: bytes):
    consulta = workbook(consulta_data)
    mapa = workbook(mapa_data)

    users_sheet = find_sheet(consulta, "Base_Perfis")
    access_sheet = find_sheet(mapa, "Botões por cargo", "Botoes por cargo")

    users = set()
    profiles = set()
    user_profile_pairs = set()
    access_rows = 0

    if users_sheet:
        for _, values, headers in iter_records(users_sheet):
            user_name = text(row_value(values, headers, "Colaborador"))
            profile_name = text(row_value(values, headers, "Grupo de acesso"))
            if user_name:
                users.add(_employee_name_identity(user_name))
            if profile_name:
                profiles.add(_profile_identity(profile_name))
            if user_name and profile_name:
                user_profile_pairs.add(
                    (_employee_name_identity(user_name), _profile_identity(profile_name))
                )

    if access_sheet:
        for _, values, headers in iter_records(access_sheet):
            profile_name = text(row_value(values, headers, "Perfil"))
            path = text(row_value(values, headers, "Caminho completo"))
            screen = text(row_value(values, headers, "Tela final"))
            button = text(row_value(values, headers, "Botão", "Botao"))
            if profile_name and (path or screen or button):
                profiles.add(_profile_identity(profile_name))
                access_rows += 1

    warnings = []
    if users_sheet is None:
        warnings.append("A aba Base_Perfis nao foi encontrada em Consulta_Acessos.xlsx.")
    if access_sheet is None:
        warnings.append("A aba Botoes por cargo nao foi encontrada em Mapa_Geral_Acessos.xlsx.")

    return {
        "arquivos": ["Consulta_Acessos.xlsx", "Mapa_Geral_Acessos.xlsx"],
        "hash_consulta": hashlib.sha256(consulta_data).hexdigest(),
        "hash_mapa": hashlib.sha256(mapa_data).hexdigest(),
        "usuarios": len(users),
        "perfis": len(profiles),
        "vinculos_usuario_perfil": len(user_profile_pairs),
        "linhas_botoes_por_cargo": access_rows,
        "abas_reconhecidas": [
            users_sheet.title if users_sheet else None,
            access_sheet.title if access_sheet else None,
        ],
        "avisos": warnings,
    }


def preview(data: bytes, filename: str):
    """Compatibilidade com a tela antiga de previa de um arquivo."""
    wb, sheets = workbook_info(data)
    return {
        "arquivo": filename,
        "hash": hashlib.sha256(data).hexdigest(),
        "abas": sheets,
        "usuarios": 0,
        "perfis": 0,
        "acessos_terminais": 0,
        "abas_reconhecidas": [ws.title for ws in wb.worksheets],
        "avisos": [
            "A nova importacao deve receber Consulta_Acessos.xlsx e "
            "Mapa_Geral_Acessos.xlsx na mesma operacao."
        ],
    }


def _load_state(session):
    employees = session.scalars(select(Funcionario)).all()
    profiles = session.scalars(select(Perfil)).all()
    accesses = session.scalars(select(Acesso)).all()

    employees_by_login = {}
    employees_by_name = defaultdict(list)
    for employee in employees:
        login_key = _employee_login_identity(employee.login)
        name_key = _employee_name_identity(employee.nome)
        if login_key:
            employees_by_login[login_key] = employee
        if name_key:
            employees_by_name[name_key].append(employee)

    profiles_by_name = {
        _profile_identity(profile.nome): profile
        for profile in profiles
        if text(profile.nome)
    }

    accesses_by_identity = {
        _access_identity(
            access.caminho_completo,
            access.tela_final,
            access.botao,
        ): access
        for access in accesses
    }

    employee_profile_links = {
        (employee_id, profile_id): link
        for employee_id, profile_id, link in session.execute(
            select(
                FuncionarioPerfil.funcionario_id,
                FuncionarioPerfil.perfil_id,
                FuncionarioPerfil,
            )
        ).all()
    }

    profile_access_links = {
        (profile_id, access_id): link
        for profile_id, access_id, link in session.execute(
            select(
                PerfilAcesso.perfil_id,
                PerfilAcesso.acesso_id,
                PerfilAcesso,
            )
        ).all()
    }

    return {
        "employees_by_login": employees_by_login,
        "employees_by_name": employees_by_name,
        "profiles_by_name": profiles_by_name,
        "accesses_by_identity": accesses_by_identity,
        "employee_profile_links": employee_profile_links,
        "profile_access_links": profile_access_links,
    }


def _get_or_create_profile(session, state, profile_name, result):
    identity = _profile_identity(profile_name)
    profile = state["profiles_by_name"].get(identity)
    if profile is not None:
        return profile

    profile = Perfil(nome=text(profile_name), situacao="Ativo")
    session.add(profile)
    session.flush([profile])
    state["profiles_by_name"][identity] = profile
    result["perfis_criados"] += 1
    return profile


def _find_employee(values, headers, state):
    login = normalize_login(
        row_value(
            values,
            headers,
            "Login",
            "Login usuário",
            "Login usuario",
            "Usuário",
            "Usuario",
            "Matrícula",
            "Matricula",
        )
    )
    if login:
        employee = state["employees_by_login"].get(login)
        if employee is not None:
            return employee

    name = text(row_value(values, headers, "Colaborador", "Nome", "Nome completo"))
    matches = state["employees_by_name"].get(_employee_name_identity(name), [])
    if len(matches) == 1:
        return matches[0]
    return None


def _import_user_profiles(session, consulta_data, state, result):
    wb = workbook(consulta_data)
    ws = find_sheet(wb, "Base_Perfis")
    if ws is None:
        raise RuntimeError("A aba Base_Perfis nao foi encontrada em Consulta_Acessos.xlsx.")

    seen_pairs = set()
    unresolved_names = set()

    for _, values, headers in iter_records(ws):
        profile_name = text(row_value(values, headers, "Grupo de acesso"))
        employee_name = text(row_value(values, headers, "Colaborador"))
        if not profile_name or not employee_name:
            continue

        profile = _get_or_create_profile(session, state, profile_name, result)
        employee = _find_employee(values, headers, state)
        if employee is None:
            unresolved_names.add(employee_name)
            continue

        pair = (employee.id, profile.id)
        if pair in seen_pairs:
            result["vinculos_repetidos_no_arquivo"] += 1
            continue
        seen_pairs.add(pair)

        if pair in state["employee_profile_links"]:
            result["vinculos_mantidos"] += 1
            continue

        link = FuncionarioPerfil(
            funcionario_id=employee.id,
            perfil_id=profile.id,
        )
        session.add(link)
        state["employee_profile_links"][pair] = link
        result["vinculos_criados"] += 1

    result["usuarios_nao_localizados"] = len(unresolved_names)
    result["usuarios_nao_localizados_nomes"] = sorted(unresolved_names)


def _get_or_create_access(session, state, path, screen, button, result):
    path = text(path)
    screen = text(screen)
    button = text(button) or "(acesso à tela)"
    identity = _access_identity(path, screen, button)
    access = state["accesses_by_identity"].get(identity)
    if access is not None:
        return access

    access = Acesso(
        caminho_completo=path,
        tela_final=screen,
        botao=button,
        chave_acesso=_natural_access_key(path, screen, button),
    )
    session.add(access)
    session.flush([access])
    state["accesses_by_identity"][identity] = access
    result["acessos_criados"] += 1
    return access


def _import_profile_accesses(session, mapa_data, state, result):
    wb = workbook(mapa_data)
    ws = find_sheet(wb, "Botões por cargo", "Botoes por cargo")
    if ws is None:
        raise RuntimeError(
            "A aba Botoes por cargo nao foi encontrada em Mapa_Geral_Acessos.xlsx."
        )

    seen_pairs = set()

    for _, values, headers in iter_records(ws):
        profile_name = text(row_value(values, headers, "Perfil"))
        path = text(row_value(values, headers, "Caminho completo"))
        screen = text(row_value(values, headers, "Tela final"))
        button = text(row_value(values, headers, "Botão", "Botao"))
        status = _status(row_value(values, headers, "Status"))

        if not profile_name or not (path or screen or button):
            continue

        profile = _get_or_create_profile(session, state, profile_name, result)
        access = _get_or_create_access(
            session,
            state,
            path,
            screen,
            button,
            result,
        )
        pair = (profile.id, access.id)

        if pair in seen_pairs:
            result["permissoes_repetidas_no_arquivo"] += 1
        seen_pairs.add(pair)

        link = state["profile_access_links"].get(pair)
        if link is None:
            link = PerfilAcesso(
                perfil_id=profile.id,
                acesso_id=access.id,
                status=status,
            )
            session.add(link)
            state["profile_access_links"][pair] = link
            result["permissoes_criadas"] += 1
        elif link.status != status:
            link.status = status
            result["permissoes_atualizadas"] += 1
        else:
            result["permissoes_mantidas"] += 1


def import_workbooks(
    session,
    consulta_data: bytes,
    mapa_data: bytes,
    consulta_filename: str = "Consulta_Acessos.xlsx",
    mapa_filename: str = "Mapa_Geral_Acessos.xlsx",
):
    """
    Mescla usuarios/grupos de Consulta_Acessos.xlsx e arvores/permissoes da aba
    'Botoes por cargo' de Mapa_Geral_Acessos.xlsx.

    A operacao e incremental: cria e atualiza o que veio, preservando registros
    ausentes. O consolidado continua derivado de FuncionarioPerfil + PerfilAcesso.
    """
    result = defaultdict(int)
    result["arquivo_usuarios"] = consulta_filename
    result["arquivo_acessos"] = mapa_filename

    state = _load_state(session)
    _import_user_profiles(session, consulta_data, state, result)
    _import_profile_accesses(session, mapa_data, state, result)
    session.flush()
    return dict(result)


def _find_companion_file(filename: str):
    """Localiza o segundo arquivo enquanto a tela antiga ainda envia um por vez."""
    project_root = Path(__file__).resolve().parent.parent
    normalized = key(filename)
    wanted = (
        "Mapa_Geral_Acessos.xlsx"
        if "consulta acessos" in normalized
        else "Consulta_Acessos.xlsx"
    )
    candidates = [
        project_root / wanted,
        project_root / "data" / wanted,
        project_root / "uploads" / wanted,
    ]
    return wanted, next((path for path in candidates if path.is_file()), None)


def import_workbook(session, data: bytes, filename: str):
    """
    Compatibilidade temporaria com a rota antiga, que ainda envia um arquivo.
    O arquivo complementar deve ficar na raiz do projeto, em data ou uploads.
    """
    normalized = key(filename)
    wanted, companion = _find_companion_file(filename)
    if companion is None:
        raise RuntimeError(
            f"A tela atual enviou somente {filename}. Coloque {wanted} na raiz "
            "do projeto e confirme novamente, ou atualize a rota para dois uploads."
        )

    companion_data = companion.read_bytes()
    if "consulta acessos" in normalized:
        return import_workbooks(
            session,
            consulta_data=data,
            mapa_data=companion_data,
            consulta_filename=filename,
            mapa_filename=companion.name,
        )
    if "mapa geral acessos" in normalized:
        return import_workbooks(
            session,
            consulta_data=companion_data,
            mapa_data=data,
            consulta_filename=companion.name,
            mapa_filename=filename,
        )
    raise RuntimeError(
        "Arquivo principal nao reconhecido. Envie Consulta_Acessos.xlsx ou "
        "Mapa_Geral_Acessos.xlsx."
    )
